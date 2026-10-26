#!/usr/bin/env python3
"""把《安装与恢复教程》.md 转成 docx。改完 md 后跑一次：python3 docs/make-tutorial-docx.py

以前是手抄一遍正文，改 md 不改脚本就会出两个版本（实际就发生过：md 2.1.4、docx 2.0.15）。
现在只解析 md，正文只有一份。
"""
import os
import re
from datetime import date

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
MD = os.path.join(HERE, "安装与恢复教程.md")
OUT = os.path.join(HERE, "安装与恢复教程.docx")
EAST = "微软雅黑"
INK = (0x23, 0x20, 0x1C)
GREY = (0x55, 0x50, 0x4B)
MONO = "Consolas"


def style_run(run, size=11, bold=False, east=EAST, ascii_font="Calibri", color=INK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    run.font.color.rgb = RGBColor(*color)


def add_inline(p, text, size=11, bold=False, color=INK):
    """**加粗** 和 `等宽` 两种行内格式，其余原样。"""
    for tok in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            style_run(p.add_run(tok[2:-2]), size, True, color=color)
        elif tok.startswith("`") and tok.endswith("`"):
            style_run(p.add_run(tok[1:-1]), size, bold, east=MONO, ascii_font=MONO, color=color)
        else:
            style_run(p.add_run(tok), size, bold, color=color)


def para(doc, text="", size=11, bold=False, space_before=0, space_after=6,
         indent=0, color=INK):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    if text:
        add_inline(p, text, size, bold, color)
    return p


def heading(doc, text, size=15):
    return para(doc, text, size=size, bold=True, space_before=16, space_after=8)


def bullet(doc, text, indent=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    add_inline(p, text, 11)


def shade(p, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    p._p.get_or_add_pPr().append(shd)


def code(doc, lines):
    for line in lines:
        p = para(doc, indent=0.6, space_after=0)
        p.paragraph_format.space_before = Pt(0)
        style_run(p.add_run(line), 10, east=MONO, ascii_font=MONO)
        shade(p, "F2F2F2")
    para(doc, "", space_after=4)


def note_box(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "FDF3EE")
    cell._tc.get_or_add_tcPr().append(shd)
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        add_inline(p, line, 10.5)


def build(doc, lines):
    i, first_heading = 0, True
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith("```"):
            j, block = i + 1, []
            while j < len(lines) and not lines[j].startswith("```"):
                block.append(lines[j].rstrip())
                j += 1
            code(doc, block)
            i = j + 1
        elif line.startswith(">"):
            block = []
            while i < len(lines) and lines[i].startswith(">"):
                block.append(lines[i].lstrip(">").strip())
                i += 1
            note_box(doc, block)
            para(doc, "", space_after=2)
        elif line.startswith("# "):
            if first_heading:
                para(doc, line[2:].strip(), size=20, bold=True, space_after=4)
                first_heading = False
            else:
                heading(doc, line[2:].strip())
            i += 1
        elif line.startswith("## "):
            heading(doc, line[3:].strip())
            i += 1
        elif line.startswith("### "):
            para(doc, line[4:].strip(), bold=True, space_before=8, space_after=2)
            i += 1
        elif line.lstrip().startswith("- "):
            indent = 0.6 if len(line) - len(line.lstrip()) >= 2 else 0
            bullet(doc, line.lstrip()[2:].strip(), indent)
            i += 1
        else:
            para(doc, line.strip(), color=GREY if first_heading else INK)
            i += 1


def main():
    with open(MD, encoding="utf-8") as fh:
        lines = fh.read().splitlines()

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    sec.left_margin = sec.right_margin = Cm(2.4)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), EAST)

    build(doc, lines)

    m = re.search(r"电视桌面 ([\d.]+)", lines[0])
    ver = m.group(1) if m else ""
    para(doc, "", space_after=2)
    para(doc, "— 电视桌面 %s · %s" % (ver, date.today().strftime("%Y-%m")),
         size=9, color=(0x88, 0x84, 0x80))

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
