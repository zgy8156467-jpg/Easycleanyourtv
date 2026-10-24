#!/usr/bin/env python3
"""开发日志 md -> docx。用法：python3 docs/make-devlog-docx.py docs/开发日志-2026-10-22.md"""
import re
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

EAST = "微软雅黑"
INK = (0x23, 0x20, 0x1C)
DIM = (0x55, 0x50, 0x4B)


def style_run(run, size=11, bold=False, east=EAST, ascii_font="Calibri", color=INK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    run.font.color.rgb = RGBColor(*color)


def inline(doc, p, text, size=11, color=INK):
    """`code` 用等宽，**bold** 加粗，其余正常。"""
    for part in re.split(r"(`[^`]*`|\*\*[^*]+\*\*)", text):
        if not part:
            continue
        run = p.add_run(part.strip("`") if part.startswith("`") else part.strip("*"))
        mono = part.startswith("`")
        style_run(run, size, bold=not mono and part.startswith("**"),
                  east="Consolas" if mono else EAST,
                  ascii_font="Consolas" if mono else "Calibri", color=color)


def para(doc, text="", size=11, bold=False, space_before=0, space_after=6, color=INK):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        inline(doc, p, text, size, color)
        if bold:
            for r in p.runs:
                r.font.bold = True
    return p


def shade_paragraph(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def main():
    src = Path(sys.argv[1] if len(sys.argv) > 1 else "docs/开发日志-2026-10-22.md")
    out = src.with_suffix(".docx")
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    sec.left_margin = sec.right_margin = Cm(2.4)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), EAST)

    for raw in src.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("# "):
            para(doc, line[2:], size=20, bold=True, space_after=4)
        elif line.startswith("## "):
            para(doc, line[3:], size=15, bold=True, space_before=16, space_after=8)
        elif line.startswith("### "):
            para(doc, line[4:], size=12, bold=True, space_before=10, space_after=4)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(4)
            inline(doc, p, line[2:])
        else:
            para(doc, line, color=DIM if line.startswith("一句话") else INK)
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
